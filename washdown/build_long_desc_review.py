"""Build the washdown long-description review spreadsheet from washdown_data.json + long_descriptions.json."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
data = json.loads((HERE / "washdown_data.json").read_text())
longs = json.loads((HERE / "long_descriptions.json").read_text())
series_names = {s["slug"]: s["title"] for s in data["series"]}
series_order = {s["slug"]: s["sortOrder"] for s in data["series"]}
products = sorted(data["products"], key=lambda p: (series_order[p["series"]], p["sortOrder"]))

FONT = "Arial"
NAVY = "1F4E9A"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
input_fill = PatternFill("solid", start_color="FFF2CC")
head_fill = PatternFill("solid", start_color=NAVY)
series_fills = ["DDEBF7", "FFFFFF"]
wrap_top = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
ws = wb.active
ws.title = "Long Descriptions"

ws["A1"] = "Washdown Guns: Long Description Review"
ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
ws["A2"] = (
    "How to review: read column E. Set Status (column H) to Approve, Edit or Reject. "
    "For Edit, type your changes or the replacement text in column I. Yellow cells are yours to fill in. "
    "Column G flags anything in the draft that isn't stated in the range PDF."
)
ws["A2"].font = Font(name=FONT, size=10, italic=True)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:I2")
ws.row_dimensions[2].height = 42

headers = ["#", "Series", "Product", "Short Description (current)", "Long Description (draft)",
           "Words", "Notes / Check", "Status", "Your Edits / Comments", "Slug (do not edit)"]
widths = [5, 16, 22, 34, 80, 8, 38, 12, 45, 24]
HR = 4
for c, (h, w) in enumerate(zip(headers, widths), start=1):
    cell = ws.cell(row=HR, column=c, value=h)
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = head_fill
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = border
    ws.column_dimensions[cell.column_letter].width = w

dv = DataValidation(type="list", formula1='"Approve,Edit,Reject"', allow_blank=True)
ws.add_data_validation(dv)

prev_series, band = None, -1
for i, p in enumerate(products, start=1):
    r = HR + i
    if p["series"] != prev_series:
        band, prev_series = band + 1, p["series"]
    entry = longs[p["slug"]]
    values = [i, series_names[p["series"]], p["name"], p["shortDescription"], entry["text"],
              f'=IF(LEN(TRIM(E{r}))=0,0,LEN(TRIM(E{r}))-LEN(SUBSTITUTE(TRIM(E{r})," ",""))+1)',
              entry["check"], None, None, p["slug"]]
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT, size=10, bold=(c == 3), color=("808080" if c == 10 else "000000"))
        cell.alignment = wrap_top
        cell.border = border
        cell.fill = input_fill if c in (8, 9) else PatternFill("solid", start_color=series_fills[band % 2])
    dv.add(ws.cell(row=r, column=8))
    ws.row_dimensions[r].height = max(75, 13 * (len(entry["text"]) // 95 + 2))

last = HR + len(products)
ws.cell(row=last + 2, column=1,
        value="Source: Nozzleworks 'Washdown Gun Range – Product line-up by series' PDF. Specs in each draft come from that PDF; "
              "anything else is listed in column G.").font = Font(name=FONT, size=9, italic=True)
ws.freeze_panes = ws.cell(row=HR + 1, column=4)
ws.auto_filter.ref = f"A{HR}:J{last}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

out = HERE / "Washdown_Long_Descriptions_Review.xlsx"
wb.save(out)
print(out)
