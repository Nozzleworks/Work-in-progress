"""Generate Quote_Template.xlsx — a quote whose totals recalculate from qty/unit price."""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLUE = "0000FF"
INPUT_FILL = PatternFill("solid", fgColor="FFF9C4")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '"$"#,##0.00;("$"#,##0.00);"-"'

FIRST_ROW, N_LINES = 17, 15
LAST_ROW = FIRST_ROW + N_LINES - 1

wb = Workbook()
ws = wb.active
ws.title = "Quote"
ws.sheet_view.showGridLines = False
for col, w in zip("ABCDEFG", [12, 44, 8, 10, 15, 16, 3]):
    ws.column_dimensions[col].width = w


def put(ref, value, bold=False, size=10, color="000000", inp=False, fmt=None, align=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, bold=bold, size=size, color=BLUE if inp else color)
    if inp:
        c.fill = INPUT_FILL
        c.border = BOX
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center")
    return c


# Header
put("A1", "QUOTATION", bold=True, size=20, color="1F3864")
put("A3", "Your Company Name", bold=True, size=12, inp=True)
put("A4", "Street address, City", inp=True)
put("A5", "Phone / Email", inp=True)
put("A6", "GST No: 000-000-000", inp=True)

put("E3", "Quote No:", bold=True, align="right"); put("F3", "Q-0001", inp=True)
put("E4", "Date:", bold=True, align="right"); put("F4", "=TODAY()", inp=True, fmt="dd/mm/yyyy")
put("E5", "Valid (days):", bold=True, align="right"); put("F5", 30, inp=True)
put("E6", "Valid Until:", bold=True, align="right"); put("F6", "=F4+F5", fmt="dd/mm/yyyy")

put("A8", "Quote To:", bold=True)
for r, txt in [(9, "Customer name"), (10, "Company"), (11, "Address"), (12, "Phone / Email")]:
    put(f"A{r}", txt, inp=True)
    ws.merge_cells(f"A{r}:B{r}")
for r in range(3, 7):
    ws.merge_cells(f"A{r}:B{r}")

put("E8", "GST Rate:", bold=True, align="right")
gst = put("F8", 0.15, inp=True, fmt="0.0%")
gst.comment = Comment("NZ GST standard rate 15% (Inland Revenue). Change here if needed.", "Template")

# Line items
headers = ["Item / Code", "Description", "Qty", "Unit", "Unit Price (excl GST)", "Line Total (excl GST)"]
for i, h in enumerate(headers):
    c = ws.cell(row=FIRST_ROW - 1, column=i + 1, value=h)
    c.font = Font(name=FONT, bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[FIRST_ROW - 1].height = 30

for r in range(FIRST_ROW, LAST_ROW + 1):
    for col, fmt in zip("ABCDE", [None, None, "#,##0.##", None, MONEY]):
        put(f"{col}{r}", None, inp=True, fmt=fmt)
    ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    put(f"F{r}", f'=IF(OR(C{r}="",E{r}=""),"",C{r}*E{r})', fmt=MONEY)
    ws[f"F{r}"].border = BOX

# Example row so the expected format is clear
ex = FIRST_ROW
ws[f"A{ex}"] = "NZ-100"
ws[f"B{ex}"] = "Example item - replace or delete this row"
ws[f"C{ex}"] = 2
ws[f"D{ex}"] = "ea"
ws[f"E{ex}"] = 125.00

units = DataValidation(type="list", formula1='"ea,hr,m,kg,set,lot"', allow_blank=True, showErrorMessage=False)
ws.add_data_validation(units)
units.add(f"D{FIRST_ROW}:D{LAST_ROW}")

# Totals
t = LAST_ROW + 2
rows = [
    (t, "Subtotal (excl GST)", f"=SUM(F{FIRST_ROW}:F{LAST_ROW})"),
    (t + 1, '="GST @ "&TEXT(F8,"0.0%")', f"=ROUND(F{t}*F8,2)"),
    (t + 2, "TOTAL (incl GST)", f"=F{t}+F{t + 1}"),
]
for r, label, formula in rows:
    ws.merge_cells(f"D{r}:E{r}")
    put(f"D{r}", label, bold=True, align="right")
    put(f"F{r}", formula, bold=True, fmt=MONEY)
    for col in "DEF":
        ws[f"{col}{r}"].border = BOX
ws[f"F{t + 2}"].font = Font(name=FONT, bold=True, size=12)
for col in "DEF":
    ws[f"{col}{t + 2}"].fill = TOTAL_FILL

# Terms
n = t + 4
put(f"A{n}", "Notes / Terms:", bold=True)
put(f"A{n + 1}", "Prices in NZD. Quote valid until the date shown above. Payment terms: 20th of the month following invoice.", inp=True)
ws.merge_cells(f"A{n + 1}:F{n + 2}")
ws[f"A{n + 1}"].alignment = Alignment(wrap_text=True, vertical="top")

# Legend
lg = n + 4
put(f"A{lg}", "How to use:", bold=True, size=9, color="7F7F7F")
put(f"A{lg + 1}", "Yellow cells with blue text are inputs - edit these. White cells are formulas - don't overwrite.", size=9, color="7F7F7F")
put(f"A{lg + 2}", "Change Qty or Unit Price and the line total, subtotal, GST and total all update automatically.", size=9, color="7F7F7F")

# Print setup
ws.print_area = f"A1:F{n + 2}"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.freeze_panes = None

wb.calculation.fullCalcOnLoad = True  # Excel computes all formulas when opened
wb.save("Quote_Template.xlsx")
