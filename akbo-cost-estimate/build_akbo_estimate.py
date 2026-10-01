"""Build AKBO_Cost_Estimate.xlsx: qty x unit price per item, with subtotal, freight and NZD conversion."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = Path(__file__).with_name("AKBO_Cost_Estimate.xlsx")

# (code, description, note) - descriptions from akbo.nl product listings via web search, Oct 2026
ITEMS = [
    ("AKRSV01-B", "Squeeze valve - multi-functional cleaning valve, blue", ""),
    ("AKRSHHL1-BL", "Low flow spray head cleaning gun", ""),
    ("AKRS001-B", "Stainless steel (AISI304) shower head, blue", ""),
    ("AKRSHHL4-BL", "Spray head cleaning gun", "Description not confirmed"),
    ("TWRWR02M", "Twister stainless steel industrial spray head, model M", "Listed twice in your request"),
    ("BMFTP11-B-PRIN", "Blue Princess economy water pistol with trigger guard", ""),
    ("AKMN001-B", "Standard flow water gun (brass), blue", ""),
    ("AKNN001-B", "Lightweight glass-reinforced PA66 water gun, blue", ""),
    ("BABTN01-B", "", "Description not confirmed"),
    ("AKRHP01-R", "Stainless steel hot-water gun with trigger protection, red", ""),
    ("AKKS003-B", "Plastic shower head, blue", ""),
    ("AKAWLU1-B", "", "Description not confirmed"),
    ("AGIL572", "Light duty water gun", ""),
    ("BABTLF1-B-V", "", "Description not confirmed"),
    ("AKWMM01", "Water saver, brass, model M, 3/4\" female", ""),
    ("WRWR02XL", "Twister spray head, model XL (?)", "Check code - probably TWRWR02XL"),
    ("TWRWR02L", "Twister stainless steel industrial spray head, model L", ""),
    ("TWRWR02M", "Twister stainless steel industrial spray head, model M", "Listed twice in your request"),
    ("SWIFB13-C17", "Stainless steel swivelling hose tail", ""),
    ("SWIFB19-C24", "Stainless steel swivelling hose tail", ""),
    ("SWIFB19-C24-34", "Stainless steel swivelling hose tail", ""),
]

F = "Arial"
bold = Font(name=F, bold=True)
norm = Font(name=F)
blue = Font(name=F, color="0000FF")
hdr_font = Font(name=F, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E78")
input_fill = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '€#,##0.00;(€#,##0.00);"-"'
NZD = '"NZ$"#,##0.00;("NZ$"#,##0.00);"-"'

wb = Workbook()
ws = wb.active
ws.title = "AKBO Estimate"

ws["A1"] = "AKBO Purchase - Cost Estimate"
ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = ("Enter prices in the yellow cells (blue text = inputs). Totals update automatically. "
            "Prices could not be pulled from akbo.nl, so they are left blank.")
ws["A2"].font = Font(name=F, italic=True, size=9)

headers = ["#", "AKBO Code", "Description", "Qty", "Unit Price (EUR, ex VAT)", "Line Total (EUR)", "Notes"]
HR = 4
for c, h in enumerate(headers, 1):
    cell = ws.cell(HR, c, h)
    cell.font, cell.fill, cell.border = hdr_font, hdr_fill, box
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

first = HR + 1
for i, (code, desc, note) in enumerate(ITEMS):
    r = first + i
    ws.cell(r, 1, i + 1).font = norm
    ws.cell(r, 2, code).font = norm
    ws.cell(r, 3, desc).font = norm
    q = ws.cell(r, 4, 1); q.font = blue; q.fill = input_fill
    p = ws.cell(r, 5); p.font = blue; p.fill = input_fill; p.number_format = EUR
    t = ws.cell(r, 6, f"=D{r}*E{r}"); t.font = norm; t.number_format = EUR
    ws.cell(r, 7, note).font = Font(name=F, italic=True, color="C00000" if note else "000000")
    for c in range(1, 8):
        ws.cell(r, c).border = box
    ws.cell(r, 1).alignment = ws.cell(r, 4).alignment = Alignment(horizontal="center")
last = first + len(ITEMS) - 1

r = last + 2
def row(label, formula=None, fmt=EUR, inp=False, bold_row=False, value=None, comment=None):
    global r
    ws.cell(r, 5, label).font = bold if bold_row else norm
    ws.cell(r, 5).alignment = Alignment(horizontal="right")
    c = ws.cell(r, 6, formula if formula is not None else value)
    c.number_format = fmt
    c.font = blue if inp else (bold if bold_row else norm)
    if inp:
        c.fill = input_fill
    c.border = box
    if comment:
        c.comment = Comment(comment, "Estimate")
    r += 1
    return r - 1

items_row = row("Items (count)", f"=COUNTA(B{first}:B{last})", fmt="0")
qty_row = row("Total qty", f"=SUM(D{first}:D{last})", fmt="0")
sub_row = row("Subtotal (EUR)", f"=SUM(F{first}:F{last})", bold_row=True)
frt_row = row("Freight / shipping (EUR)", inp=True, comment="Enter AKBO's freight quote to NZ, if known.")
tot_row = row("Total (EUR)", f"=F{sub_row}+F{frt_row}", bold_row=True)
r += 1
fx_row = row("Exchange rate (NZD per 1 EUR)", inp=True, fmt="0.0000",
             comment="Enter the current EUR to NZD rate (e.g. from your bank on the day of purchase).")
nzd_row = row("Total (NZD, before GST/duty)", f"=F{tot_row}*F{fx_row}", fmt=NZD, bold_row=True)
gst_row = row("NZ GST rate on import", value=0.15, inp=True, fmt="0%",
              comment="NZ GST 15% charged on import (on value + freight). Set to 0% to exclude.")
gst_amt = row("GST (NZD)", f"=F{nzd_row}*F{gst_row}", fmt=NZD)
grand = row("Estimated landed total (NZD incl GST)", f"=F{nzd_row}+F{gst_amt}", fmt=NZD, bold_row=True)
ws.cell(grand, 6).fill = PatternFill("solid", fgColor="DDEBF7")

r += 1
ws.cell(r, 1, "Notes").font = bold
for n in [
    "Source: product descriptions from akbo.nl product listings (AKBO Handelsmij B.V.). Prices need to be taken from your AKBO price list / dealer login.",
    "TWRWR02M appears twice in the original list, so it is included twice (qty 1 each). Delete one row if that was unintended.",
    "WRWR02XL does not match AKBO's naming; it is probably TWRWR02XL (Twister model XL). Confirm before ordering.",
    "Customs duty, import fees and brokerage are not included.",
]:
    r += 1
    ws.cell(r, 1, "- " + n).font = Font(name=F, size=9)

for col, w in zip("ABCDEFG", [5, 18, 55, 7, 22, 22, 34]):
    ws.column_dimensions[col].width = w
ws.row_dimensions[HR].height = 30
ws.freeze_panes = ws.cell(first, 1)
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.calculation.fullCalcOnLoad = True  # no cached values from openpyxl; Excel computes on open
wb.save(OUT)
print(OUT)
