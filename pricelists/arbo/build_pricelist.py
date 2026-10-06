"""Build the Nozzleworks NZ pricelist for AKBO/ARBO washdown guns.

Inputs (same folder):
  ARBO_Pricelist_washdown_guns_2026-08-15.xlsx  - supplier pricelist (EUR, qty breaks)
  akbo_conversion_table.tsv                     - export of AKBO_Part_Number_Conversion.xlsx
Output:
  Nozzleworks_Pricelist_AKBO_washdown_guns.xlsx

Weights come from AKBO_2026_Catalogue_UK.pdf (SharePoint: 01 Marketing/04 Brochures/AKBO NL).
Anything not printed in the catalogue is an ESTIMATE and is flagged amber in the workbook.
"""
import csv
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from nw_descriptions import nw_description

HERE = Path(__file__).parent
SRC = HERE / "ARBO_Pricelist_washdown_guns_2026-08-15.xlsx"
CONV = HERE / "akbo_conversion_table.tsv"
OUT = HERE / "Nozzleworks_Pricelist_AKBO_washdown_guns.xlsx"

# ---------------------------------------------------------------- weights (kg)
# (weight, source). "Cat p.N" = page number printed in AKBO 2026 catalogue.
CAT = "Catalogue"
W = {
    "AKNN001-B": (0.480, "Cat p.16/17"),
    "BMFT001-B": (0.463, "Cat p.C (Blue Princess)"),
    "BMFT001-B-PRIN": (0.463, "Cat p.C - same gun as BMFT001-B, logo only"),
    "BMFTP11-B-PRIN": (0.590, "Cat p.C"),
    **{f"KLMN001-{c}": (0.840, "Cat p.9") for c in ["B", "R", "G", "BL", "Y", "W"]},
    **{f"AKMN001-{c}": (0.835, "Cat p.11") for c in ["B", "R", "G", "BL", "Y", "W"]},
    "AKMNP01-B": (1.031, "Cat p.11"),
    "AKMNP01-B-SW12": (1.084, "Cat p.11"),
    "AKMNP01-B-SW34": (1.114, "Cat p.11"),
    "AKMN011-B": (0.706, "Cat p.6"),
    "AKMN011-B-KING": (0.706, "Cat p.6 - same gun as AKMN011-B, logo only"),
    "AKMNP11-B": (0.880, "Cat p.6"),
    "AKMNP11-B-KING": (0.880, "Cat p.6 - same gun as AKMNP11-B, logo only"),
    "AKMB001-BL": (0.840, "Cat p.20"),
    "AKMBP01-BL": (1.036, "ESTIMATE - not tabulated; AKMB001-BL + trigger guard (AKMN delta)"),
    "AKRB001-BL": (0.840, "ESTIMATE - not tabulated; taken as AKMB001-BL"),
    "AKCN001-B": (0.867, "Cat p.13"),
    "AKCNP01-B": (1.016, "Cat p.13"),
    "AKCNP01-W": (1.016, "Cat p.13"),
    "BRFT003-B": (0.480, "Cat p.C"),
    "BRFT003-B-PRIN": (0.480, "Cat p.C - same gun as BRFT003-B, logo only"),
    "BRFTP03-B-PRIN": (0.620, "Cat p.C (catalogue code BRFTP13-B-PRIN)"),
    "AKRN001-B": (0.813, "Cat p.15"),
    "AKRN001-W": (0.813, "Cat p.15"),
    "AKRNP01-B": (0.983, "Cat p.15"),
    "AKRNP01-W": (0.983, "Cat p.15"),
    "AKRHP01-R": (0.980, "Cat p.22"),
    "AKRHPB1-R": (0.980, "Cat p.22"),
    "AKRHP02-R-L40": (1.900, "Cat p.22"),
    "AKRNP01-BL-EX": (0.957, "Cat p.23"),
    "AKRN002-LB-LAT": (0.870, "Cat p.24"),
    "AKRSV01-B": (0.635, "Cat p.A"),
    "AKR001-B": (0.125, "Cat p.B (catalogue code AKRS001-B)"),
    # Baby series
    "BABTN01-B": (0.360, "Cat p.18"),
    "BABTLF1-BL": (0.365, "Cat p.37"),
    "BABTA01-B": (0.365, "ESTIMATE - not in catalogue; taken as BABTLF1-BL"),
    "BABTNAA-B": (0.370, "ESTIMATE - not in catalogue; BABTN01-B + adapters"),
    "BMBTN01-B": (0.622, "Cat p.18"),
    "BMBNN01-B": (0.650, "Cat p.37"),
    "BMBTL01-B": (0.650, "ESTIMATE - not in catalogue; taken as BMBNN01-B"),
    "BMBTA01-B": (0.650, "ESTIMATE - not in catalogue; taken as BMBNN01-B"),
    "BMBTA02-B-V": (0.650, "ESTIMATE - not in catalogue; taken as BMBNN01-B"),
    # Other water guns / spray heads
    "AGIL572": (0.200, "Cat p.29"),
    **{f"AKMSH03-{c}": (0.525, "Cat p.54 (head + brass handle)") for c in ["B", "R", "W"]},
    "AKMSHH4-BL": (0.200, "Cat p.54 (head only)"),
    **{f"AKRSH03-{c}": (0.550, "Cat p.54 (head + st.st. handle)") for c in ["B", "R", "W"]},
    "AKRSHW3-W": (0.550, "Cat p.54 (head + st.st. handle)"),
    "AKRSHH4-BL": (0.200, "Cat p.54 (head only)"),
    "AKRSHH4-W": (0.200, "Cat p.54 (head only)"),
    "AKNSH01-B": (0.245, "Cat p.56"),
    "AKMSHHL1-BL": (0.210, "ESTIMATE - AKMSHL1-B 510 g less GRIPML1-B 300 g (p.53/55)"),
    "AKMSHL1-B": (0.510, "Cat p.53"),
    "AKRSHHL1-BL": (0.210, "ESTIMATE - AKRSHL1-B 550 g less GRIPRL1-B 340 g (p.53/55)"),
    "AKRSHL1-B": (0.550, "Cat p.53"),
    "AKRSHL1-W": (0.550, "Cat p.53 (as AKRSHL1-B)"),
    "TWRWR02M": (0.345, "Cat p.57"),
    "TWRWR02L": (0.480, "Cat p.57"),
    "TWRWR02XL": (0.800, "Cat p.57"),
    "AKAWLU1-B": (0.800, "Cat p.27"),
    # Water savers / Blue Nozzle - catalogue gives no weight
    **{k: (0.300, "ESTIMATE - weight not in catalogue (p.45)") for k in
       ["AKWMS01", "AKWMM01", "AKWML01", "AKWMR01", "AKWRS01", "AKWRM01", "AKWRL01", "AKWRR01"]},
    "BWNM045": (0.250, "ESTIMATE - weight not in catalogue (p.46)"),
    # Shower heads
    "AKKS003-B": (0.260, "Cat p.26"),
    "AKKS003-W": (0.260, "Cat p.26"),
    "AKKS003-B-19": (0.300, "ESTIMATE - 260 g shower (p.26) + MESC019 hose tail est."),
    "AKKS003-W-19": (0.300, "ESTIMATE - 260 g shower (p.26) + MESC019 hose tail est."),
    # Ball valves - no weight in catalogue
    "BALR001-B": (0.400, "ESTIMATE - weight not in catalogue (p.38)"),
    "BALRCP1-B": (0.550, "ESTIMATE - weight not in catalogue (p.38)"),
    "BALRCP1-NB": (0.550, "ESTIMATE - weight not in catalogue (p.38)"),
    "BARRCP2-B": (0.600, "ESTIMATE - weight not in catalogue (p.38)"),
}
# Swivels & hose tails: catalogue (p.68-71) gives dimensions only, no weights.
_EST_FIT = "ESTIMATE - catalogue gives dimensions only (p.68-71)"
for pre in ["SWM", "SWC", "SWR"]:
    for sfx, kg in [("09-12", .090), ("09-34", .100), ("14-34", .130), ("14-34X34", .150),
                    ("18-1X1", .250), ("09-12MX12M", .090), ("09-12Mx12M", .090), ("09-12MX34M", .100)]:
        W[f"{pre}{sfx}"] = (kg, _EST_FIT)
for k, kg in {
    "SWIFM1/2": .060, "SWIFR1/2": .060, "SWIFB10": .080,
    "SWIFB13-C17": .090, "SWIFB13-C17-SP": .130, "SWIFB13-C20": .090, "SWIFB13-C20-SP": .130,
    "SWIFB16-C21": .100, "SWIFB16-C21-SP": .140, "SWIFB19-C24": .120, "SWIFB19-C24-34": .130,
    "SWIFB19-C24-34S": .170, "SWIFB19-C24-SP": .160,
    "MESS013": .030, "MESS019": .040, "MESS025": .060,
    "MESC013": .035, "MESC016": .040, "MESC019": .045,
    "RVST013-C17": .040, "RVST013-C17-SP": .080, "RVST013-C20": .040, "RVST013-C20-12F": .050,
    "RVST013-C20-34F": .060, "RVST013-C20-SP": .080, "RVST016-C21": .050, "RVST016-C21-SP": .090,
    "RVST019-C24": .060, "RVST019-C24-SP": .100, "RVST025": .060,
    "RVSW13C17-12C": .060, "RVSW13C17-34F": .070, "RVSW13C20-12C": .060, "RVSW13C20-12F": .060,
    "RVSW13C20-12FS": .100, "RVSW13C20-34F": .070, "RVSW16C21-12C": .070, "RVSW16C21-12CS": .110,
    "RVSW16C21-12F": .070, "RVSW16C21-34F": .080, "RVSW19C24-34C": .090, "RVSW19C24-34F": .090,
    "RVSW19C24-1C": .110, "RVSW19C24-1F": .110, "RVSW25C36-1C": .130,
}.items():
    W[k] = (kg, _EST_FIT)

# Pricelist codes whose spelling differs from the catalogue/conversion table.
CODE_NOTES = {
    "BRFTP03-B-PRIN": "Pricelist code BRFTP03 vs catalogue/conversion BRFTP13 - confirm with AKBO.",
    "AKR001-B": "Catalogue code is AKRS001-B - confirm with AKBO.",
    "AKMBP01-BL": "Catalogue only lists AKMBP01-B-SWM14-3/4 (with swivel) - confirm model.",
}

# ---------------------------------------------------------------- styles
FONT = "Arial"
f_body = Font(name=FONT, size=9)
f_bold = Font(name=FONT, size=9, bold=True)
f_input = Font(name=FONT, size=9, color="0000FF")
f_link = Font(name=FONT, size=9, color="008000")
f_title = Font(name=FONT, size=14, bold=True)
f_section = Font(name=FONT, size=11, bold=True, color="FFFFFF")
f_hdr = Font(name=FONT, size=9, bold=True, color="FFFFFF")
fill_section = PatternFill("solid", fgColor="4A4A4A")
fill_hdr = PatternFill("solid", fgColor="666666")
fill_group = PatternFill("solid", fgColor="E7E6E6")
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_amber = PatternFill("solid", fgColor="FFE699")
fill_red = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
EUR = '€#,##0.00;(€#,##0.00);"-"'
NZD = '$#,##0.00;($#,##0.00);"-"'
KG = '0.000'

# ---------------------------------------------------------------- read sources
conv_rows = list(csv.reader(CONV.open(encoding="utf-8"), delimiter="\t"))
src = openpyxl.load_workbook(SRC).active

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Pricelist"

# ---------------------------------------------------------------- inputs block
ws["A1"] = "NOZZLEWORKS PRICELIST - AKBO / ARBO WASHDOWN GUNS & FITTINGS"
ws["A1"].font = f_title
ws["A2"] = (f"Built from AKBO price list \"{src['A1'].value}\". "
            "Item order and quantity breaks follow AKBO's list. Prices NZD, GST exclusive.")
ws["A2"].font = Font(name=FONT, size=9, italic=True)

inputs = [
    ("Exchange rate (EUR per 1 NZD)", 0.46, "0.000", "Per Andrew 28-Sep-2026 - to be reviewed. NZD cost = EUR / rate."),
    ("Freight rate (NZD per kg)", 30, '$#,##0.00', "Per Andrew 28-Sep-2026 - placeholder, to be updated."),
    ("Volumetric weight factor (% on net weight)", 0.25, "0.0%", "Per Andrew 30-Sep-2026 - freight kg = net weight x (1 + factor)."),
    ("Tax / duty allowance (% of product cost)", 0.01, "0.0%", "Per Andrew 28-Sep-2026 - 1% on NZD product cost."),
    ("List price markup (% on landed cost)", 3.5, "0%", "Per Andrew 30-Sep-2026 - 350% default, to be reviewed. List = landed x (1 + markup)."),
]
r = 4
ws.cell(r, 1, "INPUTS (edit yellow cells)").font = f_bold
for i, (label, val, fmt, note) in enumerate(inputs, start=r + 1):
    ws.cell(i, 1, label).font = f_body
    c = ws.cell(i, 3, val)
    c.font, c.fill, c.number_format, c.border = f_input, fill_input, fmt, box
    ws.cell(i, 4, note).font = Font(name=FONT, size=8, italic=True, color="666666")
FX, FRT, VOL, DUTY, MU = "$C$5", "$C$6", "$C$7", "$C$8", "$C$9"

ws["A10"] = ("Landed cost / pc = EUR price / FX x (1 + duty) + net weight kg x (1 + volumetric factor) x freight rate.   "
             "List price / pc = landed cost x (1 + markup %).")
ws["A10"].font = Font(name=FONT, size=8, italic=True)
ws["A11"] = ("Legend:  red Nozzleworks Part No = no NW part number assigned yet   |   "
             "amber weight = ESTIMATE, not printed in AKBO 2026 catalogue   |   "
             "blank price = AKBO gives no price at that break")
ws["A11"].font = Font(name=FONT, size=8, italic=True)

# ---------------------------------------------------------------- column headers
HDR = 13
heads = ["AKBO Product code", "Nozzleworks Part No", "AKBO description", "Proposed NW description", "Net weight (kg)", "Weight source",
         "Supplier price EUR / pc", None, None,
         "Landed cost NZD / pc", None, None,
         "List price NZD / pc (ex GST)", None, None,
         "Flags / notes"]
for col, h in enumerate(heads, start=1):
    c = ws.cell(HDR, col, h)
    c.font, c.fill, c.border = f_hdr, fill_section, box
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for c1, c2 in [(7, 9), (10, 12), (13, 15)]:
    ws.merge_cells(start_row=HDR, start_column=c1, end_row=HDR, end_column=c2)
ws.row_dimensions[HDR].height = 30
ws.freeze_panes = ws.cell(HDR + 1, 5)

# ---------------------------------------------------------------- body
def is_tier_row(row):
    # header rows have tier labels in C:E and no description in B
    return row[1] is None and any(isinstance(v, str) and "pcs" in v for v in row[2:5])


out = HDR + 1
blank_run = 0
item_rows, missing, estimates = [], [], []
src_rows = [[c.value for c in row] for row in src.iter_rows(min_row=3, max_col=7)]
for idx, row in enumerate(src_rows):
    a, b, c, d, e, f, g = row
    if all(v is None for v in row):
        blank_run += 1
        if blank_run == 1:
            out += 1
        continue
    blank_run = 0
    if a == "Product code":
        continue  # our own column header replaces it
    if is_tier_row(row):
        tiers = [v for v in (c, d, e)]
        for base in (7, 10, 13):
            for k, t in enumerate(tiers):
                cell = ws.cell(out, base + k, t)
                cell.font, cell.fill, cell.border = f_hdr, fill_hdr, box
                cell.alignment = Alignment(horizontal="center")
        label = a if isinstance(a, str) else "Qty break"
        cell = ws.cell(out, 1, label)
        cell.font = f_hdr
        for col in range(1, 7):
            ws.cell(out, col).fill = fill_hdr
        ws.cell(out, 16).fill = fill_hdr
        out += 1
        continue
    if isinstance(a, str) and b is None and not isinstance(c, (int, float)):
        # section / sub-group heading (ignore AKBO's internal discount % cells)
        nxt = next((r for r in src_rows[idx + 1:] if any(v is not None for v in r)), None)
        main = nxt is not None and nxt[0] == "Product code"
        for col in range(1, 17):
            ws.cell(out, col).fill = fill_section if main else fill_group
        cell = ws.cell(out, 1, a)
        cell.font = f_section if main else f_bold
        out += 1
        continue
    if a is None:
        continue  # AKBO's discount-% helper rows
    # ---- item row
    code = str(a).strip()
    desc = " ".join(str(b).split())
    rr = out
    ws.cell(rr, 1, code).font = f_bold
    ws.cell(rr, 2, f'=IFERROR(INDEX(Conversion!$C$2:$C$500,MATCH($A{rr},Conversion!$B$2:$B$500,0)),"NOT ASSIGNED")').font = f_body
    ws.cell(rr, 3, desc).font = f_body
    nwd = nw_description(code, desc)
    ws.cell(rr, 4, nwd).font = f_bold if nwd else f_body
    wkg, wsrc = W.get(code, (None, "MISSING - add weight"))
    cw = ws.cell(rr, 5, wkg)
    cw.number_format, cw.font = KG, f_input
    ws.cell(rr, 6, wsrc).font = Font(name=FONT, size=8, color="666666")
    if wsrc.startswith("ESTIMATE") or wkg is None:
        cw.fill = fill_amber
        estimates.append((code, wkg, wsrc))
    notes = []
    for k, v in enumerate((c, d, e)):
        col = 7 + k
        if isinstance(v, (int, float)):
            pc = ws.cell(rr, col, round(v, 6))
            pc.number_format, pc.font = EUR, f_input
        elif isinstance(v, str):
            notes.append(f"AKBO note at 3rd break: '{v}'")
        L = 10 + k
        S = 13 + k
        pcol = openpyxl.utils.get_column_letter(col)
        lcol = openpyxl.utils.get_column_letter(L)
        ws.cell(rr, L, f'=IF({pcol}{rr}="","",{pcol}{rr}/{FX}*(1+{DUTY})+$E{rr}*(1+{VOL})*{FRT})').number_format = NZD
        ws.cell(rr, S, f'=IF({lcol}{rr}="","",ROUND({lcol}{rr}*(1+{MU}),2))').number_format = NZD
        ws.cell(rr, L).font = f_body
        ws.cell(rr, S).font = f_bold
    if code in CODE_NOTES:
        notes.append(CODE_NOTES[code])
    static = " ".join(notes)
    static_q = static.replace('"', '""')
    ws.cell(rr, 16, (
        f'=TRIM(IF($B{rr}="NOT ASSIGNED","No NW part no. ","")'
        f'&IFERROR(IF(INDEX(Conversion!$J$2:$J$500,MATCH($A{rr},Conversion!$B$2:$B$500,0))="Needs review",'
        f'"Conversion table: needs review. ",""),"")'
        f'&IF(LEFT($F{rr},8)="ESTIMATE","Weight estimated. ","")&"{static_q}")'
    )).font = Font(name=FONT, size=8)
    for col in range(1, 17):
        ws.cell(rr, col).border = box
    item_rows.append(rr)
    out += 1

last = out - 1
# red fill for missing NW part numbers
ws.conditional_formatting.add(
    f"B{HDR+1}:B{last}",
    FormulaRule(formula=[f'$B{HDR+1}="NOT ASSIGNED"'], fill=fill_red, font=Font(name=FONT, size=9, bold=True, color="C00000")))

widths = {"A": 18, "B": 20, "C": 50, "D": 55, "E": 9, "F": 30, "G": 10, "H": 10, "I": 10,
          "J": 11, "K": 11, "L": 11, "M": 12, "N": 12, "O": 12, "P": 48}
for k, v in widths.items():
    ws.column_dimensions[k].width = v
for rr in item_rows:
    ws.cell(rr, 3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(rr, 4).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(rr, 6).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(rr, 16).alignment = Alignment(wrap_text=True, vertical="top")

ws["M5"] = "Items listed"
ws["N5"] = f'=COUNT(G{HDR+1}:G{last})'
ws["M6"] = "No NW part no."
ws["N6"] = f'=COUNTIF(B{HDR+1}:B{last},"NOT ASSIGNED")'
ws["M7"] = "Estimated weights"
ws["N7"] = f'=COUNTIF(F{HDR+1}:F{last},"ESTIMATE*")'
for rr in (5, 6, 7):
    ws.cell(rr, 13).font = f_bold
    ws.cell(rr, 14).font = f_body
ws.print_title_rows = f"{HDR}:{HDR}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- Conversion sheet
cs = wb.create_sheet("Conversion")
for i, row in enumerate(conv_rows, start=1):
    for j, v in enumerate(row, start=1):
        cell = cs.cell(i, j, v if v != "" else None)
        cell.font = f_hdr if i == 1 else f_body
        if i == 1:
            cell.fill = fill_section
    if i > 1 and row[9] == "Needs review":
        for j in range(1, 12):
            cs.cell(i, j).fill = fill_amber
note_r = len(conv_rows) + 2
cs.cell(note_r, 1, "Copy of AKBO_Part_Number_Conversion.xlsx (SharePoint: 01 Marketing/04 Brochures/AKBO NL), "
                   "values as at 28-Sep-2026. Add new AKBO -> Nozzleworks codes here; the Pricelist picks them up.").font = \
    Font(name=FONT, size=8, italic=True)
for col, wdt in zip("ABCDEFGHIJK", [26, 20, 22, 30, 36, 26, 30, 50, 40, 13, 60]):
    cs.column_dimensions[col].width = wdt
cs.freeze_panes = "C2"

# ---------------------------------------------------------------- Flags sheet
fs = wb.create_sheet("Flags")
fs["A1"] = "Items to action"
fs["A1"].font = f_title
fs["A3"] = "1. AKBO pricelist items with no Nozzleworks part number in the conversion table"
fs["A3"].font = f_bold
conv_codes = {r[1] for r in conv_rows[1:]}
src_codes = [ws.cell(rr, 1).value for rr in item_rows]
no_pn = [(c, ws.cell(rr, 3).value) for c, rr in zip(src_codes, item_rows) if c not in conv_codes]
r = 4
for h, col in (("AKBO code", 1), ("Description", 2), ("Nozzleworks Part No (live)", 3)):
    fs.cell(r, col, h).font = f_hdr
    fs.cell(r, col).fill = fill_hdr
for code, desc in no_pn:
    r += 1
    fs.cell(r, 1, code).font = f_body
    fs.cell(r, 2, desc).font = f_body
    fs.cell(r, 3, f'=IFERROR(INDEX(Conversion!$C$2:$C$500,MATCH($A{r},Conversion!$B$2:$B$500,0)),"NOT ASSIGNED")').font = f_body
fs.conditional_formatting.add(f"C5:C{r}", FormulaRule(formula=['$C5="NOT ASSIGNED"'], fill=fill_red))
r += 1
fs.cell(r, 1, f"{len(no_pn)} items at build time (column C updates as codes are added to the Conversion sheet).").font = \
    Font(name=FONT, size=8, italic=True)

r += 2
fs.cell(r, 1, "2. Weights not printed in the AKBO 2026 catalogue (estimated - please confirm with AKBO)").font = f_bold
r += 1
for h, col in (("AKBO code", 1), ("Estimated kg", 2), ("Basis", 3)):
    fs.cell(r, col, h).font = f_hdr
    fs.cell(r, col).fill = fill_hdr
for code, kg, basis in estimates:
    r += 1
    fs.cell(r, 1, code).font = f_body
    fs.cell(r, 2, kg).number_format = KG
    fs.cell(r, 2).font = f_body
    fs.cell(r, 3, basis).font = f_body

r += 2
fs.cell(r, 1, "3. Part-number / data queries").font = f_bold
for code, msg in CODE_NOTES.items():
    r += 1
    fs.cell(r, 1, code).font = f_body
    fs.cell(r, 3, msg).font = f_body
for cr in conv_rows[1:]:
    if cr[9] == "Needs review" and cr[1] in src_codes:
        r += 1
        fs.cell(r, 1, cr[1]).font = f_body
        fs.cell(r, 3, "Conversion table: " + cr[10]).font = f_body
r += 1
fs.cell(r, 1, "AGIL572 / AKAWLU1-B").font = f_body
fs.cell(r, 3, "AKBO show only one price plus a pack-qty note (48 pcs EUR 9.87 / 20 pcs EUR 66.81) - kept as note, not priced.").font = f_body
r += 1
fs.cell(r, 1, "Heavy duty spray guns").font = f_body
fs.cell(r, 3, "AKBO give no >=50 pcs price for most guns - left blank.").font = f_body
fs.column_dimensions["A"].width = 22
fs.column_dimensions["B"].width = 60
fs.column_dimensions["C"].width = 90

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print(f"items={len(item_rows)} no_pn={len(no_pn)} estimates={len(estimates)} last_row={last}")
print("missing weights:", [c for c, k, s in estimates if k is None])
